from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from docx import Document


NOME_BASE = "Embrapii_seleção_analista_2026_questao04_Base_Automação.xlsx"
NOME_MOLDE = "Embrapii_seleção_analista_2026_questao04_Molde_Ficha_Projeto (1).docx"

ROTULOS_PARA_COLUNAS = {
    "Código do Projeto": "Projeto_cód",
    "Empresa Contratante (cód.)": "Empresas Contratantes_cód",
    "Status do Projeto": "Projeto_status",
    "Tipo de Projeto": "Tipo de Projeto",
    "Data do Contrato": "Data_Contrato",
    "Data de Início": "Data_Início",
    "Data de Término": "Data_Término",
    "Duração (dias)": "Data_Duração_dias",
    "Código da Unidade": "Unidade Embrapii_cód",
    "Status da Unidade": "Unidade Embrapii_status",
    "Região": "Unidade Embrapii_Região",
    "UF": "Unidade Embrapii_UF",
    "Tipo de Instituição": "Unidade Embrapii_Tipo de Instituição",
    "Fonte de Recursos": "Fin_Fonte de Recursos",
    "Valor Aporte Embrapii (R$)": "Fin_Valor aporte Embrapii",
    "Valor Contrapartida Empresa (R$)": "Fin_Valor Contrapartida Empresa",
    "Valor Aporte Unidade (R$)": "Fin_Valor Aporte Unidade",
    "Valor Total (R$)": "Fin_Valor total",
    "TRL Inicial": "TRL Inicial",
    "TRL Final": "TRL Final",
    "Entregável": "Entregável",
    "Área de Aplicação": "Classificação Área de Aplicação_nível e subnível",
    "Tecnologia Habilitadora": "Classificação_Tecnologia Habilitadora_nível e subnível",
    "Tecnologia Verde": "Classificação_Tecnologia Verde_nível e subnível",
    "Número de Pedidos de PI": "Número de Pedidos de PI",
}


def localizar_arquivo(caminho_informado: Path | None, candidatos: list[Path]) -> Path:
    if caminho_informado:
        if caminho_informado.exists():
            return caminho_informado
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho_informado}")

    for candidato in candidatos:
        if candidato.exists():
            return candidato

    nomes = "\n".join(f"- {candidato}" for candidato in candidatos)
    raise FileNotFoundError(f"Nenhum dos arquivos esperados foi encontrado:\n{nomes}")


def valor_para_texto(valor: object) -> str:
    if pd.isna(valor):
        return ""
    return str(valor)


def definir_texto_celula(celula, texto: str) -> None:
    paragrafo = celula.paragraphs[0]

    if not paragrafo.runs:
        paragrafo.add_run()

    primeiro_run = paragrafo.runs[0]
    primeiro_run.text = ""

    for run in paragrafo.runs[1:]:
        run.text = ""

    partes = texto.split("\n")
    primeiro_run.text = partes[0]
    for parte in partes[1:]:
        primeiro_run.add_break()
        primeiro_run.add_text(parte)

    for paragrafo_extra in celula.paragraphs[1:]:
        for run in paragrafo_extra.runs:
            run.text = ""


def preencher_ficha(caminho_molde: Path, projeto: pd.Series) -> Document:
    documento = Document(caminho_molde)

    for tabela in documento.tables:
        for linha in tabela.rows:
            if len(linha.cells) < 2:
                continue

            rotulo = linha.cells[0].text.strip()
            coluna = ROTULOS_PARA_COLUNAS.get(rotulo)
            if coluna is None:
                continue

            if coluna not in projeto.index:
                raise KeyError(f"Coluna ausente na planilha: {coluna}")

            definir_texto_celula(linha.cells[1], valor_para_texto(projeto[coluna]))

    return documento


def componente_nome_arquivo(valor: object) -> str:
    texto = valor_para_texto(valor).strip()
    texto = re.sub(r'[<>:"/\\|?*]+', "_", texto)
    return texto or "sem_codigo"


def gerar_fichas(caminho_base: Path, caminho_molde: Path, pasta_saida: Path) -> list[Path]:
    projetos = pd.read_excel(caminho_base, engine="openpyxl").head(5)
    pasta_saida.mkdir(parents=True, exist_ok=True)

    arquivos_gerados = []
    for _, projeto in projetos.iterrows():
        codigo = componente_nome_arquivo(projeto["Projeto_cód"])
        documento = preencher_ficha(caminho_molde, projeto)
        caminho_saida = pasta_saida / f"Ficha_Projeto_{codigo}.docx"
        documento.save(caminho_saida)
        arquivos_gerados.append(caminho_saida)

    return arquivos_gerados


def criar_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Gera fichas de projeto EMBRAPII em DOCX a partir de uma planilha."
    )
    parser.add_argument("--base", type=Path, help="Caminho da planilha .xlsx.")
    parser.add_argument("--molde", type=Path, help="Caminho do molde .docx.")
    parser.add_argument(
        "--saida",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="Pasta onde as fichas geradas serão salvas.",
    )
    return parser


def main() -> None:
    pasta_script = Path(__file__).resolve().parent
    pasta_downloads = Path.home() / "Downloads"
    args = criar_parser().parse_args()

    caminho_base = localizar_arquivo(
        args.base,
        [
            pasta_script / NOME_BASE,
            pasta_script / "Base_projetos_embrapii.xlsx",
            pasta_downloads / NOME_BASE,
            pasta_downloads / "Base_projetos_embrapii.xlsx",
        ],
    )
    caminho_molde = localizar_arquivo(
        args.molde,
        [
            pasta_script / NOME_MOLDE,
            pasta_script / "Molde_Ficha_Projeto_EMBRAPII.docx",
            pasta_downloads / NOME_MOLDE,
            pasta_downloads / "Molde_Ficha_Projeto_EMBRAPII.docx",
        ],
    )

    arquivos = gerar_fichas(caminho_base, caminho_molde, args.saida)
    for arquivo in arquivos:
        print(arquivo)


if __name__ == "__main__":
    main()
